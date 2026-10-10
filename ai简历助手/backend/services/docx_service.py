"""
Word 文档导出服务

将优化后的 HTML 简历转换为 .docx 格式。
"""

import io
import logging
import re

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

logger = logging.getLogger(__name__)


def html_to_docx(html_content: str, title: str = "简历") -> io.BytesIO:
    """将 HTML 简历内容转换为 DOCX 文档

    简化解析：提取文本内容并按段落/标题结构写入 Word 文档。

    Args:
        html_content: 优化后的 HTML 简历内容
        title: 文档标题

    Returns:
        包含 DOCX 文件内容的 BytesIO 对象
    """
    doc = Document()

    # 设置默认字体
    style = doc.styles["Normal"]
    font = style.font
    font.name = "宋体"
    font.size = Pt(11)

    # 设置页边距
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    # 简单的 HTML 文本提取和结构化
    # 移除 script/style 标签
    clean_html = re.sub(r"<script[^>]*>.*?</script>", "", html_content, flags=re.DOTALL)
    clean_html = re.sub(r"<style[^>]*>.*?</style>", "", clean_html, flags=re.DOTALL)

    # 提取标题 (h1-h3)
    headings = re.findall(r"<h[1-3][^>]*>(.*?)</h[1-3]>", clean_html, re.DOTALL)

    # 按段落分割
    # 先处理 <br> 和 </p> 为换行
    text_content = re.sub(r"<br\s*/?>", "\n", clean_html)
    text_content = re.sub(r"</p>", "\n", text_content)
    text_content = re.sub(r"</li>", "\n", text_content)
    text_content = re.sub(r"</h[1-3]>", "\n", text_content)

    # 移除所有 HTML 标签
    text_content = re.sub(r"<[^>]+>", "", text_content)

    # 清理多余空白
    text_content = re.sub(r"\n\s*\n", "\n\n", text_content)
    text_content = text_content.strip()

    # 解析内容并写入文档
    lines = text_content.split("\n")
    heading_set = set(h.strip() for h in headings)

    for line in lines:
        line = line.strip()
        if not line:
            continue

        # 检查是否是标题行
        if line in heading_set:
            doc.add_heading(line, level=2)
        elif re.match(r"^(个人信息|教育背景|工作经历|项目经历|专业技能|自我评价|求职意向)", line):
            doc.add_heading(line, level=2)
        elif line.startswith(("•", "·", "-", "–", "—")):
            # 列表项
            p = doc.add_paragraph(line[1:].strip(), style="List Bullet")
        else:
            doc.add_paragraph(line)

    # 保存到 BytesIO
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)

    return output
